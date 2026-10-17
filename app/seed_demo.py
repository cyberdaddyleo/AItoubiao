"""演示数据：python -m app.seed_demo

创建一个示例主体（示例消防工程有限公司）、资料库、订阅，并用 tests/fixtures 里
2026-10-15 采集的真实公告样本生成标讯与一个进行中的投标项目，方便第一次打开就能看到完整流程。
示例主体与资料均为虚构，正式使用前可在「主体与订阅」里删除。
"""
from __future__ import annotations

import io
import shutil
import uuid
from pathlib import Path

from . import db, matching, writer
from .config import BASE_DIR, FILES_DIR, now_str
from .crawler import ccgp, runner
from .security import ensure_admin

FIX = BASE_DIR / "tests" / "fixtures"


def _tender_docx() -> bytes:
    from docx import Document

    d = Document()
    d.add_heading("第一章 招标公告", 1)
    d.add_paragraph("项目名称：吉林省消防救援总队-2026年消防专项建设项目（第五标段）")
    d.add_paragraph("项目编号：JLSZC20260089720")
    d.add_paragraph("提交投标文件截止时间：2026年11月05日 09点30分（北京时间）")
    d.add_paragraph("开标时间：2026年11月05日 09点30分（北京时间）")
    d.add_paragraph("投标保证金：人民币伍万元整（￥50000元），须于2026年11月04日17:00前到账。")
    d.add_paragraph("对招标文件有疑问的，应于2026年10月28日17:00前以书面形式提出澄清要求。")
    d.add_paragraph("投标有效期：90日历天")
    d.add_heading("第二章 投标人资格要求", 1)
    for x in [
        "1. 满足《中华人民共和国政府采购法》第二十二条规定；",
        "2. 具备消防设施工程专业承包二级及以上资质；",
        "3. 具有ISO9001质量管理体系认证证书；",
        "4. 具有有效的安全生产许可证；",
        "5. 拟派项目经理须具备机电工程专业二级及以上注册建造师执业资格；",
        "6. 近三年具有类似消防设备供货安装业绩。",
    ]:
        d.add_paragraph(x)
    d.add_heading("第三章 评标办法（综合评分法）", 1)
    rows = [
        ["评审因素", "分值", "评分标准"],
        ["投标报价", "30", "满足招标文件要求且投标价格最低的为评标基准价，得满分"],
        ["设备技术参数响应", "20", "完全满足技术参数得 20 分，每有一项负偏离扣 2 分，★号条款负偏离按无效投标处理"],
        ["供货及安装实施方案", "15", "方案完整、针对性强、措施可行的得 11-15 分；基本可行 6-10 分；较差 0-5 分"],
        ["质量保证措施", "10", "质量目标明确、保证措施完善的得 7-10 分"],
        ["售后服务与培训方案", "10", "响应时间、备品备件、培训计划完善的得 7-10 分"],
        ["类似项目业绩", "10", "每提供一个类似业绩得 5 分，最高 10 分"],
        ["企业信誉与认证", "5", "具有质量、环境、职业健康安全管理体系认证的每项得 1-2 分"],
    ]
    t = d.add_table(rows=len(rows), cols=3)
    for i, row in enumerate(rows):
        for j, v in enumerate(row):
            t.cell(i, j).text = v
    d.add_heading("第四章 投标人须知", 1)
    d.add_paragraph("★ 投标人须承诺供货及安装工期不超过90日历天，否则投标无效。")
    d.add_paragraph("★ 消防设备须提供国家消防产品强制性认证证书（CCCF），否则投标无效。")
    d.add_paragraph("投标文件未按规定加盖公章或签字的，作无效投标处理。")
    d.add_paragraph("投标报价超过最高限价的，其投标将被否决。")
    d.add_paragraph("电子投标文件须使用 CA 数字证书加密后上传至吉林省公共资源交易一体化平台。")
    d.add_heading("第五章 投标文件格式", 1)
    for x in ["一、投标函", "二、开标一览表", "三、法定代表人身份证明", "四、授权委托书", "五、资格证明文件", "六、类似项目业绩", "七、拟派项目人员", "八、供货及安装实施方案", "九、技术偏离表", "十、中小企业声明函"]:
        d.add_paragraph(x)
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def run() -> None:
    db.init_db()
    ensure_admin()
    runner.ensure_default_sources()
    if db.q1("SELECT id FROM companies WHERE name = '示例消防工程有限公司'"):
        print("演示数据已存在")
        return
    t = now_str()
    cid = db.insert(
        "companies",
        {
            "name": "示例消防工程有限公司",
            "short_name": "示例消防",
            "credit_code": "91530100MA6DEMO01X",
            "legal_person": "张示例",
            "reg_capital": "3000万元",
            "founded_date": "2012-06-18",
            "address": "云南省昆明市五华区示例路 88 号",
            "postcode": "650000",
            "phone": "0871-60000000",
            "email": "bid@example.com",
            "contact_person": "李投标",
            "contact_phone": "13800000000",
            "business_scope": "消防设施工程设计与施工；消防设备销售、安装与维保；安防工程。",
            "intro": "本公司为演示用虚构企业。",
            "is_sme": 1,
            "staff_count": 86,
            "annual_revenue_wan": 5200,
            "industries": "消防工程, 消防设备, 安防",
            "created_at": t,
            "updated_at": t,
        },
    )

    def lib(cat, name, **kw):
        data = kw.pop("data", {})
        db.insert(
            "library_items",
            {
                "company_id": cid,
                "category": cat,
                "name": name,
                "number": kw.get("number", ""),
                "issuer": kw.get("issuer", ""),
                "level": kw.get("level", ""),
                "valid_from": kw.get("valid_from", ""),
                "valid_to": kw.get("valid_to", ""),
                "tags": kw.get("tags", ""),
                "content": kw.get("content", ""),
                "data_json": db.jdump(data),
                "file_ids": "[]",
                "created_at": t,
                "updated_at": t,
            },
        )

    lib("basic", "营业执照", number="91530100MA6DEMO01X", issuer="昆明市市场监督管理局")
    lib("basic", "2025年度财务审计报告", issuer="示例会计师事务所")
    lib("basic", "近6个月纳税证明")
    lib("basic", "近6个月社保缴纳证明")
    lib("basic", "信用中国查询截图", valid_to="2026-11-30")
    lib("qualification", "消防设施工程专业承包资质证书", level="二级", number="D253000001", issuer="云南省住房和城乡建设厅", valid_to="2028-12-31")
    lib("qualification", "安全生产许可证", number="（云）JZ安许证字〔2024〕000001", issuer="云南省住房和城乡建设厅", valid_to="2027-05-31")
    lib("qualification", "ISO9001质量管理体系认证证书", number="00125Q0001", issuer="示例认证有限公司", valid_to="2026-10-20")
    lib("qualification", "ISO14001环境管理体系认证证书", number="00125E0001", issuer="示例认证有限公司", valid_to="2027-08-15")
    lib("performance", "昆明市某三甲医院消防系统改造工程", tags="消防, 医院", content="火灾自动报警系统、喷淋系统改造及调试。", data={"client": "某三甲医院", "amount_wan": 320, "sign_date": "2025-03-01", "industry": "消防"})
    lib("performance", "曲靖市某中学消防设施维保项目", tags="消防, 维保, 学校", content="三年期消防设施维护保养。", data={"client": "某中学", "amount_wan": 86, "sign_date": "2024-09-10", "industry": "消防维保"})
    lib("performance", "某物流园区安防监控工程", tags="安防", data={"client": "某物流公司", "amount_wan": 150, "sign_date": "2023-11-20", "industry": "安防"})
    lib("personnel", "王项目", number="云153000001", valid_to="2027-12-31", data={"position": "项目经理", "certs": "二级注册建造师（机电工程）", "title": "工程师"})
    lib("personnel", "赵技术", data={"position": "技术负责人", "certs": "一级注册消防工程师", "title": "高级工程师"})
    lib(
        "solution",
        "消防设备供货安装实施方案（通用素材）",
        tags="消防, 供货, 安装",
        content=(
            "# 实施组织\n成立项目部，设项目经理、技术负责人、安全员、质检员和材料员，实行项目经理负责制。\n\n"
            "# 供货保障\n设备到货前与采购人确认交货地点和时间；到货后会同采购人开箱验收，核对型号、数量、合格证和 3C 认证证书。\n\n"
            "# 安装调试\n严格按设计图纸和《火灾自动报警系统施工及验收标准》施工；隐蔽工程验收合格后方可进入下道工序；系统联动调试合格后提交调试报告。\n\n"
            "# 进度保障\n编制进度计划横道图，关键节点周报；材料提前备货，必要时增加作业班组。\n\n"
            "# 安全文明施工\n设专职安全员，动火作业办理审批，现场设置警示标志，做到工完场清。"
        ),
    )
    lib(
        "solution",
        "售后服务与培训方案（通用素材）",
        tags="售后, 培训, 维保",
        content="# 服务响应\n设立 7×24 小时服务电话，接到报修后在承诺时间内到场。\n\n# 备品备件\n常用易损件本地备货。\n\n# 培训\n对使用单位值班人员进行操作培训和应急演练，培训后发放操作手册。",
    )
    db.insert(
        "subscriptions",
        {"company_id": cid, "name": "消防工程与设备", "keywords_include": "消防, 火灾报警, 喷淋", "keywords_exclude": "食堂, 食材, 物业", "notice_types": "tender", "enabled": 1, "created_at": t, "updated_at": t},
    )
    db.insert(
        "subscriptions",
        {"company_id": cid, "name": "医院设备维保", "keywords_include": "维保, 维护保养", "purchaser_keywords": "医院", "notice_types": "tender", "enabled": 1, "created_at": t, "updated_at": t},
    )

    # 用真实公告样本生成标讯
    src = db.q1("SELECT * FROM sources WHERE type = 'ccgp_list'")
    if FIX.exists():
        items = ccgp.parse_list((FIX / "ccgp_list_gkzb.html").read_text(encoding="utf-8"), "http://www.ccgp.gov.cn/cggg/zygg/gkzb/")
        items += ccgp.parse_list((FIX / "ccgp_list_dfgg.html").read_text(encoding="utf-8"), "http://www.ccgp.gov.cn/cggg/dfgg/")
        detail_html = (FIX / "ccgp_detail_tender.html").read_text(encoding="utf-8")
        award_html = (FIX / "ccgp_detail_award.html").read_text(encoding="utf-8")
        target = None
        for it in items:
            nid = runner.store_item(it, src)
            if not nid:
                continue
            if "t20261015_27496413" in it["url"]:
                runner.apply_detail(nid, ccgp.parse_detail(detail_html, it["url"]), detail_html)
                target = nid
            elif "27496512" in it["url"]:
                runner.apply_detail(nid, ccgp.parse_detail(award_html, it["url"]), award_html)
            matching.match_notice(nid)
        if target:
            pid = db.insert(
                "projects",
                {
                    "company_id": cid,
                    "notice_id": target,
                    "name": "吉林省消防救援总队-2026年消防专项建设项目（第五标段）",
                    "project_no": "JLSZC20260089720",
                    "purchaser": "吉林省消防救援总队",
                    "agency": "吉林省公共资源交易中心（吉林省政府采购中心）",
                    "region": "长春市",
                    "category": "货物/设备/政法、消防、检测设备/消防设备",
                    "budget": 2550000,
                    "status": "preparing",
                    "signup_deadline": "2026-10-22",
                    "deadline": "2026-11-05 09:30",
                    "open_time": "2026-11-05 09:30",
                    "bond_amount": 50000,
                    "owner_id": 1,
                    "created_by": 1,
                    "created_at": t,
                    "updated_at": t,
                },
            )
            db.ex("UPDATE matches SET status = 'project' WHERE notice_id = ? AND company_id = ?", (target, cid))
            db.insert("project_events", {"project_id": pid, "user_id": 1, "action": "立项", "comment": "演示数据", "created_at": t})
            # 招标文件
            sub = FILES_DIR / "demo"
            sub.mkdir(parents=True, exist_ok=True)
            path = sub / (uuid.uuid4().hex[:8] + "_招标文件（演示）.docx")
            path.write_bytes(_tender_docx())
            fid = db.insert("files", {"filename": "招标文件（演示）.docx", "stored_path": str(path), "mime": "", "size": path.stat().st_size, "kind": "tender", "created_at": t})
            from .routers.projects import _add_tender_doc

            _add_tender_doc(pid, fid, use_ai=False)
            did = writer.create_bid_doc(pid, 1)
            writer.update_price(did, {"items": [
                {"name": "火灾自动报警控制器（联动型）", "spec": "示例品牌 JB-QB-XX", "qty": 2, "unit": "台", "unit_price": 186000},
                {"name": "智能疏散指示系统", "spec": "示例品牌 A 型", "qty": 1, "unit": "套", "unit_price": 920000},
                {"name": "安装调试及培训", "spec": "", "qty": 1, "unit": "项", "unit_price": 1080000},
            ], "delivery": "合同签订后 90 日历天内完成供货安装", "warranty": "3 年", "tax_note": "含税、含运输安装调试培训等全部费用"}, 1)
            for s in db.q("SELECT id FROM bid_sections WHERE doc_id = ? AND kind = 'tech' ORDER BY order_no LIMIT 2", (did,)):
                writer.generate_section(s["id"], user_id=1, use_ai=False)
    print("演示数据已创建：示例消防工程有限公司")


if __name__ == "__main__":
    run()
