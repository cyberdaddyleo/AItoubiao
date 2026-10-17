"""端到端：采集 → 匹配 → 立项 → 招标文件解析 → 标书 → 体检 → 导出 → 审核 → 投递 → 结果关联。

网络请求用 2026-10-15 取下的真实页面样本代替；AI 用假模型代替（验证接线与约束）。
"""
import io
import json

import pytest
from conftest import fixture_text
from docx import Document
from fastapi.testclient import TestClient

from app import llm
from app.crawler import base as crawler_base
from app.main import app


@pytest.fixture(scope="module")
def client():
    with TestClient(app) as c:
        r = c.post("/api/auth/login", json={"username": "admin", "password": "Admin12345"})
        assert r.status_code == 200, r.text
        yield c


@pytest.fixture(autouse=True)
def fake_network(monkeypatch):
    list_html = fixture_text("ccgp_list_gkzb.html")
    detail_html = fixture_text("ccgp_detail_tender.html")

    def fake_get(self, url, *, interval=None, check_block=True):
        self.request_count += 1
        if url.endswith("/gkzb/"):
            return list_html
        if "index_" in url:
            return list_html
        return detail_html

    monkeypatch.setattr(crawler_base.Fetcher, "get", fake_get)


def make_tender_docx() -> bytes:
    d = Document()
    d.add_heading("第一章 招标公告", 1)
    d.add_paragraph("项目编号：JLSZC20260089720")
    d.add_paragraph("提交投标文件截止时间：2026年11月05日 09点30分（北京时间）")
    d.add_paragraph("投标保证金：人民币伍万元整（￥50000元），须于2026年11月04日17:00前到账。")
    d.add_paragraph("投标有效期：90日历天")
    d.add_heading("第二章 投标人资格要求", 1)
    d.add_paragraph("1. 具备消防设施工程专业承包二级及以上资质；")
    d.add_paragraph("2. 具有ISO9001质量管理体系认证证书；")
    d.add_paragraph("3. 具有安全生产许可证；")
    d.add_paragraph("4. 近三年具有类似消防工程业绩。")
    d.add_heading("第三章 评标办法", 1)
    t = d.add_table(rows=5, cols=3)
    for i, row in enumerate([["评审因素", "分值", "评分标准"], ["投标报价", "30", "满足要求且价格最低得满分"], ["施工方案", "25", "方案完整、针对性强得 20-25 分"], ["质量保证措施", "15", "措施完善得 15 分"], ["类似业绩", "10", "每个业绩得 5 分，最高 10 分"]]):
        for j, v in enumerate(row):
            t.cell(i, j).text = v
    d.add_heading("第四章 投标须知", 1)
    d.add_paragraph("★ 投标人须承诺工期不超过90日历天，否则投标无效。")
    d.add_paragraph("投标文件未按规定加盖公章的，作无效投标处理。")
    d.add_heading("第五章 投标文件格式", 1)
    for x in ["一、投标函", "二、开标一览表", "三、法定代表人身份证明", "四、授权委托书", "五、资格证明文件", "六、施工组织设计", "七、技术偏离表"]:
        d.add_paragraph(x)
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def test_full_flow(client, monkeypatch):
    # ---------- 主体与资料库
    comp = client.post(
        "/api/companies",
        json={
            "name": "云南星衡消防工程有限公司",
            "short_name": "星衡消防",
            "credit_code": "91530000MA6TEST01X",
            "legal_person": "张三",
            "reg_capital": "2000万元",
            "address": "云南省昆明市五华区测试路1号",
            "phone": "0871-60000000",
            "email": "bid@example.com",
            "contact_person": "李四",
            "is_sme": 1,
            "staff_count": 86,
        },
    ).json()
    cid = comp["id"]
    other = client.post("/api/companies", json={"name": "云南另一家安防有限公司", "short_name": "另一家"}).json()

    def lib(cat, **kw):
        r = client.post("/api/library", json={"company_id": cid, "category": cat, **kw})
        assert r.status_code == 200, r.text
        return r.json()

    lib("basic", name="营业执照", number="91530000MA6TEST01X")
    lib("basic", name="2025年度财务审计报告")
    lib("basic", name="近6个月纳税证明")
    lib("basic", name="近6个月社保缴纳证明")
    q1 = lib("qualification", name="消防设施工程专业承包资质证书", level="二级", number="D253000001", valid_to="2028-12-31")
    iso = lib("qualification", name="ISO9001质量管理体系认证证书", number="00125Q0001", valid_to="2026-10-20")
    lib("performance", name="昆明市某医院消防改造工程", data={"client": "某医院", "amount_wan": 320, "sign_date": "2025-03-01"}, tags="消防")
    lib("personnel", name="王五", number="云153000001", data={"position": "项目经理", "certs": "一级建造师（机电）"})
    lib("solution", name="消防工程施工方案模板", content="# 施工准备\n组织技术交底，编制专项施工方案。\n\n# 质量控制\n严格执行三检制度，隐蔽工程验收合格后方可进入下道工序。\n\n# 安全文明施工\n设置专职安全员，动火作业办理审批。")

    # ---------- 订阅 & 采集
    sub = client.post("/api/subscriptions", json={"company_id": cid, "name": "消防工程", "keywords_include": "消防", "notice_types": "tender"}).json()
    assert sub["id"]
    sources = client.get("/api/sources").json()["sources"]
    ccgp_src = [s for s in sources if s["type"] == "ccgp_list"][0]
    client.put(f"/api/sources/{ccgp_src['id']}", json={"config": {"channels": ["zygg/gkzb"], "max_pages": 1, "interval": 1}})
    from app.crawler import runner

    stats = runner.run_source(ccgp_src["id"], trigger="manual")
    assert stats["status"] == "ok", stats
    assert stats["new"] == 14
    # 标题含「消防」的 4 条会抓详情
    assert stats["detail"] == 4
    # 第二次运行：全部已采集，增量停止
    stats2 = runner.run_source(ccgp_src["id"], trigger="manual")
    assert stats2["new"] == 0

    lst = client.get("/api/notices", params={"company_id": cid}).json()
    assert lst["total"] >= 4
    top = lst["items"][0]
    assert top["match"]["level"] == "能投"
    assert any("消防" in r for r in top["match"]["reasons"])
    nid = top["id"]
    detail = client.get(f"/api/notices/{nid}").json()
    assert detail["project_no"] == "JLSZC20260089720"
    assert detail["budget"] == 2550000.0

    # ---------- 立项
    proj = client.post("/api/projects", json={"company_id": cid, "notice_id": nid}).json()
    pid = proj["id"]
    assert proj["deadline"] == "2026-11-05 09:30"
    assert proj["status"] == "preparing"

    # ---------- 招标文件解析
    up = client.post("/api/files", files={"file": ("招标文件.docx", make_tender_docx(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}).json()
    td = client.post(f"/api/projects/{pid}/tender_docs", json={"file_id": up["id"]}).json()
    a = td["analysis"]
    assert len(a["qualifications"]) >= 4
    items = {s["item"]: s for s in a["scoring"]}
    assert items["施工方案"]["weight"] == 25 and items["施工方案"]["category"] == "技术"
    assert items["投标报价"]["category"] == "价格"
    assert any("★" in s["text"] for s in a["starred"])
    assert "投标函" in a["format"]["toc"]
    assert a["key_info"]["validity"] == "90日历天"
    ck = client.get(f"/api/projects/{pid}/checklist").json()["items"]
    status_of = {i["requirement"].split("（")[0]: i["status"] for i in ck if i["group"] == "资格要求"}
    assert any(v == "已具备" for k, v in status_of.items() if "消防设施工程" in k)
    assert any(v == "已过期" for k, v in status_of.items() if "ISO9001" in k)
    assert any(v == "缺失" for k, v in status_of.items() if "安全生产许可证" in k)

    # ---------- 标书
    doc = client.post(f"/api/projects/{pid}/bid_doc").json()
    titles = [s["title"] for s in doc["sections"]]
    assert "投标函" in titles and "施工组织设计" in titles and "开标一览表" in titles
    assert any(t.startswith("施工方案（25分）") for t in titles)
    letter = [s for s in doc["sections"] if s["form_key"] == "bid_letter"][0]
    assert "云南星衡消防工程有限公司" in letter["content"]
    assert "JLSZC20260089720" in letter["content"]
    qual = [s for s in doc["sections"] if s["kind"] == "qualification"][0]
    assert "消防设施工程专业承包资质证书" in qual["content"]

    # 报价超预算 → 体检必须报错
    did = doc["id"]
    price = client.put(f"/api/bid_docs/{did}/price", json={"items": [{"name": "消防设备及安装", "spec": "详见清单", "qty": 1, "unit": "项", "unit_price": 2600000}]}).json()
    assert price["total"] == 2600000
    rep = client.post(f"/api/bid_docs/{did}/check", json={}).json()
    cats = {(i["category"], i["level"]) for i in rep["items"]}
    assert ("报价", "error") in cats, rep["items"]
    assert ("资质有效期", "error") in cats
    assert ("未完成内容", "error") in cats
    assert not rep["passed"]

    # 修正报价；把另一个主体的名字写进技术章节 → 主体混用
    client.put(f"/api/bid_docs/{did}/price", json={"items": [{"name": "消防设备及安装", "qty": 1, "unit": "项", "unit_price": 2480000}]})
    doc = client.get(f"/api/projects/{pid}/bid_doc").json()
    letter = [s for s in doc["sections"] if s["form_key"] == "bid_letter"][0]
    assert "贰佰肆拾捌万元整" in letter["content"]
    tech = [s for s in doc["sections"] if s["kind"] == "tech"][0]
    client.put(f"/api/sections/{tech['id']}", json={"content": "本方案参考了云南另一家安防有限公司的做法。"})
    rep = client.post(f"/api/bid_docs/{did}/check", json={}).json()
    assert any(i["category"] == "主体混用" for i in rep["items"])
    assert not any(i["category"] == "报价" and "超过预算" in i["title"] for i in rep["items"])

    # 版本记录：覆盖已有内容时保留旧版本
    client.put(f"/api/sections/{tech['id']}", json={"content": "本方案参考了云南另一家安防有限公司的做法。第二版。"})
    vers = client.get(f"/api/sections/{tech['id']}/versions").json()
    assert len(vers) == 1 and "第二版" not in vers[0]["preview"]

    # ---------- 导出 Word
    r = client.get(f"/api/bid_docs/{did}/export")
    assert r.status_code == 200
    d = Document(io.BytesIO(r.content))
    assert d.core_properties.author == "云南星衡消防工程有限公司"
    text = "\n".join(p.text for p in d.paragraphs)
    assert "投 标 文 件" in text and "投标函" in text
    assert len(d.tables) >= 2
    r = client.get(f"/api/projects/{pid}/checklist/export")
    assert r.status_code == 200

    # ---------- 审核流
    r = client.post(f"/api/projects/{pid}/review", json={"action": "submit"})
    assert r.status_code == 200
    r = client.post(f"/api/projects/{pid}/review", json={"action": "approve"})
    assert r.status_code == 400 and "必须修改" in r.json()["detail"]
    r = client.post(f"/api/projects/{pid}/review", json={"action": "approve", "force": True})
    assert r.json()["status"] == "ready"
    r = client.post(f"/api/projects/{pid}/submission", json={"platform_name": "吉林省公共资源交易一体化平台", "ca_holder": "李四", "submit_method": "electronic"})
    assert r.json()["status"] == "submitted"

    # ---------- 串标风险：另一个主体也立项同一公告
    client.post("/api/library", json={"company_id": other["id"], "category": "basic", "name": "营业执照"})
    p2 = client.post("/api/projects", json={"company_id": other["id"], "notice_id": nid}).json()
    rep = client.post(f"/api/bid_docs/{did}/check", json={}).json()
    assert any(i["category"] == "串标风险" for i in rep["items"])
    client.delete(f"/api/projects/{p2['id']}")

    # ---------- 中标结果公告自动关联
    award_text = (
        "一、项目编号：JLSZC20260089720\n二、项目名称：吉林省消防救援总队-2026年消防专项建设项目\n"
        "三、中标（成交）信息\n中标（成交）供应商名称：云南星衡消防工程有限公司\n中标金额：248.000000万元\n"
        "采购人信息\n名 称：吉林省消防救援总队\n"
    )
    res = client.post("/api/notices/import", json={"title": "吉林省消防救援总队-2026年消防专项建设项目中标(成交)结果公告", "text": award_text}).json()
    pj = client.get(f"/api/projects/{pid}").json()
    assert pj["result_notice_id"] == res["id"]
    assert "云南星衡消防工程有限公司" in pj["winner"]
    notes = client.get("/api/notifications").json()["items"]
    assert any("疑似中标" in n["title"] for n in notes)
    r = client.post(f"/api/projects/{pid}/result", json={"result": "won", "result_amount": 2480000})
    assert r.json()["status"] == "won"
    hist = client.get("/api/library", params={"company_id": cid, "category": "history_bid"}).json()
    assert any("（中标）" in h["name"] for h in hist)

    # ---------- 看板
    dash = client.get("/api/stats/dashboard").json()
    assert dash["notices_7d"] >= 15
    ov = client.get("/api/stats/overview").json()
    assert ov["won"] == 1

    # ---------- 不保存平台密码
    r = client.post("/api/platform_accounts", json={"company_id": cid, "platform_name": "某平台", "password": "x"})
    assert r.status_code == 400


def test_ai_paths(client, monkeypatch):
    """用假模型验证 AI 接线：资格复核、招标文件 AI 解析、技术章节起草的约束。"""
    from app import db

    db.set_setting("llm_provider", "openai")
    db.set_setting("llm_base_url", "http://localhost:9/v1")
    db.set_setting("llm_model", "fake")
    calls = []

    def fake_chat(prompt, *, system="", feature="", max_tokens=None, temperature=None, retries=2, cfg=None):
        calls.append(feature)
        if feature == "qualification":
            return json.dumps({"level": "差一项", "met": ["具备消防资质"], "missing": ["安全生产许可证"], "uncertain": [], "summary": "缺安全生产许可证"}, ensure_ascii=False)
        if feature.startswith("tender_"):
            key = feature[len("tender_"):]
            payload = {
                "qualifications": {"qualifications": [{"text": "具备消防设施工程专业承包二级资质", "type": "资质", "source": "第2页"}]},
                "scoring": {"scoring": [{"item": "施工方案", "weight": 25, "criteria": "方案完整", "category": "技术", "source": "第3页"}]},
                "reject_clauses": {"reject_clauses": [{"text": "未盖章无效", "starred": False, "tip": "逐页盖章", "source": "第4页"}]},
                "format": {"format": {"toc": ["投标函", "施工组织设计"], "required_forms": ["投标函"], "rules": []}},
                "key_info": {"key_info": {"deadline": "2026-11-05 09:30", "bond_amount": 50000}, "clarify_questions": ["工期起算点不明确"], "summary": "消防专项工程"},
            }[key]
            return json.dumps(payload, ensure_ascii=False)
        if feature == "write_tech":
            return "```markdown\n# 施工部署\n本项目工期【待补充：工期】。参考云南另一家安防有限公司经验。\n```"
        return "{}"

    monkeypatch.setattr(llm, "chat", fake_chat)
    assert llm.enabled()

    lst = client.get("/api/notices").json()
    mid = lst["items"][0]["match"]["id"]
    m = client.post(f"/api/matches/{mid}/ai_check").json()
    assert m["level"] == "差一项" and m["level_by"] == "ai"
    assert "安全生产许可证" in m["missing"]

    projects = client.get("/api/projects").json()["items"]
    pid = projects[0]["id"]
    td = client.get(f"/api/projects/{pid}/tender_docs").json()[0]
    from app import tender_analysis

    text = open(db.get("tender_docs", td["id"])["text_path"], encoding="utf-8").read()
    merged = tender_analysis.analyze(text, [], use_ai=True)
    assert merged["by"] == "ai"
    assert merged["clarify_questions"] == ["工期起算点不明确"]
    assert merged["key_info"]["validity"] == "90日历天"  # 规则结果保留，AI 结果补充

    doc = client.get(f"/api/projects/{pid}/bid_doc").json()
    tech = [s for s in doc["sections"] if s["kind"] == "tech"][1]
    from app import writer

    out = writer.generate_section(tech["id"], use_ai=True)
    assert out["ai_generated"] == 1
    assert "```" not in out["content"]
    # 其他主体名称被替换成占位，不会进入标书
    assert "云南另一家安防有限公司" not in out["content"]
    assert "【待补充" in out["content"]
    db.set_setting("llm_provider", "none")
